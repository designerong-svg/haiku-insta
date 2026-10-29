"""data/calendar.json → 편집 가능한 엑셀(data/haiku_calendar_365.xlsx)

시트: 안내 / 캘린더 / 데이목록 / 요약(COUNTIF 수식)
'배정된 시 ID'·'시 배정 상태' 열(노란색)이 후니가 채워 나가는 입력 칸입니다.
"""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "haiku_calendar_365.xlsx"
FONT = "Pretendard"
SEASON_KO = {"spring": "봄", "summer": "여름", "autumn": "가을", "winter": "겨울"}
SEASON_FILL = {"봄": "FBE3E6", "여름": "DDF1E4", "가을": "FCE8CF", "겨울": "E1E9F2"}
INK, WHITE, YELLOW, LIGHT_YELLOW = "181616", "FFFFFF", "FFF2A8", "FFF9D6"

rows = json.loads((ROOT / "data" / "calendar.json").read_text("utf-8"))
f = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)
thin = Side(style="thin", color="D9D9D9")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)

wb = Workbook()

# ───────────── 캘린더 ─────────────
ws = wb.active
ws.title = "캘린더"
cols = [  # (헤더, 폭, 값 함수)
    ("일차", 6, lambda r: r["day_no"]),
    ("날짜", 12, lambda r: dt.date.fromisoformat(r["date"])),
    ("요일", 5, lambda r: r["weekday"]),
    ("주차", 6, lambda r: r["week_no"]),
    ("계절", 6, lambda r: SEASON_KO[r["season"]]),
    ("현재 절기", 9, lambda r: r["solar_term_period"]),
    ("오늘의 절기", 10, lambda r: r["solar_term_today"] or None),
    ("오늘의 데이·기념일", 22, lambda r: r["event_names"] or None),
    ("데이 종류", 16, lambda r: r["event_kinds"] or None),
    ("오늘의 테마", 28, lambda r: r["primary_theme"]),
    ("일상 상황", 20, lambda r: r["situation"]),
    ("인생 시기", 16, lambda r: r["life_stage"]),
    ("톤", 10, lambda r: r["tone"]),
    ("계절 이미지", 28, lambda r: r["seasonal_imagery"]),
    ("공감 포인트/주의", 40, lambda r: r["inclusion_note"] or None),
    ("날짜 확인 방식", 24, lambda r: r["source_check"] or None),
    ("계절 시 대체 가능", 11, lambda r: "O" if r["fallback_to_season_ok"] else "X"),
    ("배정된 시 ID", 12, lambda r: r.get("poem_id")),
    ("시 배정 상태", 12, lambda r: r["poem_status"]),
]
for c, (h, w, _) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = f(bold=True, color=WHITE)
    cell.fill = fill(YELLOW if c >= 18 else INK)
    if c >= 18:
        cell.font = f(bold=True, color=INK)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = box
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 30

for i, r in enumerate(rows, 2):
    has_event = bool(r["event_names"])
    for c, (h, w, fn) in enumerate(cols, 1):
        cell = ws.cell(row=i, column=c, value=fn(r))
        cell.font = f()
        cell.border = box
        cell.alignment = Alignment(vertical="top", wrap_text=True,
                                   horizontal="center" if w <= 12 else "left")
        if c == 2:
            cell.number_format = "yyyy-mm-dd"
        if c == 3 and r["weekday"] in ("토", "일"):
            cell.font = f(color="C0392B" if r["weekday"] == "일" else "2E5C9E", bold=True)
        if c == 5:
            cell.fill = fill(SEASON_FILL[cell.value])
        if c in (8, 10) and has_event:
            cell.font = f(bold=True)
            cell.fill = fill(LIGHT_YELLOW) if c == 8 else PatternFill()
        if c >= 18:
            cell.fill = fill(LIGHT_YELLOW)
    ws.row_dimensions[i].height = 42 if r["inclusion_note"] else 28

last = len(rows) + 1
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
dv = DataValidation(type="list", formula1='"unassigned,draft,approved,posted"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"S2:S{last}")

# ───────────── 데이목록 ─────────────
wd = wb.create_sheet("데이목록")
heads = [("날짜", 12), ("요일", 5), ("데이·기념일", 26), ("종류", 18), ("테마", 30), ("톤", 10),
         ("공감 포인트/주의", 46), ("날짜 확인 방식", 30)]
for c, (h, w) in enumerate(heads, 1):
    cell = wd.cell(row=1, column=c, value=h)
    cell.font = f(bold=True, color=WHITE)
    cell.fill = fill(INK)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = box
    wd.column_dimensions[get_column_letter(c)].width = w
n = 2
for r in rows:
    if not r["event_names"]:
        continue
    vals = [dt.date.fromisoformat(r["date"]), r["weekday"], r["event_names"], r["event_kinds"],
            r["primary_theme"], r["tone"], r["inclusion_note"] or None, r["source_check"] or None]
    for c, v in enumerate(vals, 1):
        cell = wd.cell(row=n, column=c, value=v)
        cell.font = f()
        cell.border = box
        cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if c <= 2 else "left")
        if c == 1:
            cell.number_format = "yyyy-mm-dd"
        if "미확인" in str(vals[7] or ""):
            cell.fill = fill("FFE0E0")
    wd.row_dimensions[n].height = 30
    n += 1
wd.freeze_panes = "A2"
wd.auto_filter.ref = f"A1:H{n - 1}"

# ───────────── 요약 (수식) ─────────────
wsum = wb.create_sheet("요약")
wsum.column_dimensions["A"].width = 26
wsum.column_dimensions["B"].width = 12
wsum.column_dimensions["C"].width = 44
R = f"$2:$%d" % last


def rng(col):
    return f"캘린더!${col}$2:${col}${last}"


def put(r, label, formula, note=None, bold=False, pct=False):
    a = wsum.cell(row=r, column=1, value=label)
    b = wsum.cell(row=r, column=2, value=formula)
    a.font = f(bold=bold)
    b.font = f(bold=bold)
    b.alignment = Alignment(horizontal="center")
    if pct:
        b.number_format = "0.0%"
    if note:
        wsum.cell(row=r, column=3, value=note).font = f(color="7F7F7F", size=9)


def section(r, title):
    c = wsum.cell(row=r, column=1, value=title)
    c.font = f(bold=True, color=WHITE)
    for col in (1, 2, 3):
        wsum.cell(row=r, column=col).fill = fill(INK)


r = 1
section(r, "전체")
put(2, "총 일수", f"=COUNTA({rng('B')})", "2026-10-30 ~ 2027-10-29")
put(3, "데이·기념일이 있는 날", f"=COUNTA({rng('H')})")
put(4, "데이 없는 날(계절·상황·인생 시기 테마)", "=B2-B3")
put(5, "계절 시 대체 불가(X) 날", f'=COUNTIF({rng("Q")},"X")', "절기 당일 — 해당 절기 시를 우선 배정")

section(7, "시 배정 진행률 (R·S열을 채우면 자동 갱신)")
put(8, "approved (게시 가능)", f'=COUNTIF({rng("S")},"approved")')
put(9, "draft (검수 대기)", f'=COUNTIF({rng("S")},"draft")')
put(10, "unassigned (미배정)", f'=COUNTIF({rng("S")},"unassigned")')
put(11, "approved 비율", "=B8/B2", pct=True, bold=True)

section(13, "계절별 일수")
for k, s in enumerate(("봄", "여름", "가을", "겨울")):
    put(14 + k, s, f'=COUNTIF({rng("E")},A{14 + k})')

section(19, "데이 종류별 일수 (한 날에 여러 종류면 각각 집계)")
for k, s in enumerate(("명절", "법정·사회기념일", "절기", "14일 데이", "숫자·MZ 데이", "생활")):
    put(20 + k, s, f'=COUNTIF({rng("I")},"*"&A{20 + k}&"*")')

section(27, "인생 시기별 일수")
stages = sorted({r["life_stage"] for r in rows})
for k, s in enumerate(stages):
    put(28 + k, s, f'=COUNTIF({rng("L")},A{28 + k})')
put(28 + len(stages), "합계", f"=SUM(B28:B{27 + len(stages)})", "총 일수와 같아야 함(검산)", bold=True)
for row in wsum.iter_rows(min_row=1, max_row=40):
    for c in row:
        if c.font.name != FONT:
            c.font = f()

# ───────────── 안내 ─────────────
wg = wb.create_sheet("안내", 0)
wg.column_dimensions["A"].width = 22
wg.column_dimensions["B"].width = 100
lines = [
    ("하이쿠 365일 테마 캘린더", None, "title"),
    ("기간", "2026-10-30 ~ 2027-10-29 (365일). 하이쿠 채널 '오늘의 하이쿠' 자동 게시용 테마 계획표", None),
    ("", "", None),
    ("여섯 겹 구조", None, "h"),
    ("① 계절·24절기", "절기 당일은 해당 절기 시를 우선, 평일엔 '현재 절기'의 계절 이미지(계절 이미지 열)로 시를 고른다", None),
    ("② 일상 상황", "요일별 순환(월 출근길 · 화 나른한 오후 · 수 퇴근길 · 목 따뜻한 차 · 금 금요일 밤 · 토 주말 아침 · 일 잠 못 드는 밤). 장마철(6/25~7/25) 수·일요일은 '비 오는 날'", None),
    ("③ 인생 시기", "주 단위로 17가지 인생 시기를 돌린다(위로·새 출발·번아웃·이별·중년의 고독·나이 듦·부모와 자식·사랑·우정 등). 수능·새해·설·새 학기·가정의 달·추석·노인의 날 주간은 날짜에 맞게 고정", None),
    ("④ 14일 데이", "다이어리·밸런타인·화이트·블랙·로즈·키스·실버·뮤직(그린)·포토·와인·무비·허그데이", None),
    ("⑤ 숫자·MZ 데이", "삼겹살(3/3)·오이(5/2)·빼빼로(11/11)·세계 고양이의 날(8/8) 등. 비공식 마케팅 데이는 가볍게, 맞는 시가 없으면 계절 시", None),
    ("⑥ 명절·가족·사회 기념일", "설·정월대보름·단오·칠석·추석·복날, 어린이날·어버이날·스승의 날·성년의 날·부부의 날·장애인의 날·노인의 날·한글날 등", None),
    ("", "", None),
    ("선정 원칙", None, "h"),
    ("누구도 소외시키지 않기", "연인·솔로·가족 없음·부모 부재 등 어느 쪽에도 상처가 되지 않는 날만 넣고, 연인 데이는 가족·친구·나 자신에게 건네는 마음으로 푼다", None),
    ("제외한 날", "정치적 해석이 갈릴 수 있는 날(4.16, 5.18 등)과 성별·이념 논쟁이 붙는 날(세계 여성의 날 등)은 일단 제외. 필요하면 추가", None),
    ("시가 안 맞으면", "'계절 시 대체 가능'이 O인 날은 억지로 끼우지 말고 계절 시로 둔다", None),
    ("", "", None),
    ("사용법 (직접 채우는 칸)", None, "h"),
    ("노란색 R·S열", "R열 '배정된 시 ID'에 data/queue.json 의 시 ID, S열 '시 배정 상태'에 상태를 적는다. 요약 시트의 진행률이 자동으로 바뀐다", None),
    ("입력 예시", "R: h001  /  S: approved      (상태 값: unassigned → draft → approved → posted)", None),
    ("권장 진행", "한 달 분량씩 원문·작가를 확인하며 채운다. 퍼블릭 도메인 고전 작가만 사용하고, AI가 만든 원문은 俳句データベース 등으로 확인한 뒤 approved", None),
    ("", "", None),
    ("날짜 확인 상태 (솔직한 메모)", None, "h"),
    ("계산으로 구한 날짜", "24절기(태양 황경 15° 단위, KST), 설날·정월대보름·부처님오신날·단오·칠석·추석(음양력 변환), 복날(경일 계산). 2027 설날은 한국 시간 기준 2/7", None),
    ("검색으로 확인", "부부의 날(5/21), 수능(2026-11-19), 14일 데이 목록, 삼겹살·빼빼로·오이데이", None),
    ("통상 고정일(재확인 권장)", "식목일·어린이날·어버이날·스승의 날·장애인의 날·노인의 날·한글날 등 법정기념일. 행정안전부 목록 페이지는 열리지 않아 이번엔 대조하지 못함", None),
    ("확인 필요(분홍색)", "4/2 '사이데이'는 후니님이 언급한 비공식 데이로 이름·유래 미확인. 8/14는 출처마다 뮤직데이/그린데이로 갈림. 12/14는 허그데이(머니데이로 부르는 곳도 있음)", None),
]
for i, (a, b, kind) in enumerate(lines, 1):
    ca = wg.cell(row=i, column=1, value=a)
    cb = wg.cell(row=i, column=2, value=b)
    ca.font, cb.font = f(bold=True), f()
    ca.alignment = Alignment(vertical="top", wrap_text=True)
    cb.alignment = Alignment(vertical="top", wrap_text=True)
    if kind == "title":
        ca.font = f(bold=True, size=16)
        wg.row_dimensions[i].height = 28
    elif kind == "h":
        for c in (ca, cb):
            c.fill = fill(INK)
        ca.font = f(bold=True, color=WHITE)
    elif b and len(b) > 90:
        wg.row_dimensions[i].height = 30
wg.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT, "rows", len(rows), "events", n - 2)
